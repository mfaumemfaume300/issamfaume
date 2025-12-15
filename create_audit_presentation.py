#!/usr/bin/env python3
"""
Script to create a professional PowerPoint presentation on Audit Completion & Reporting
with 35 slides featuring modern design and cool color scheme.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE


def create_presentation():
    """Create a comprehensive audit completion and reporting presentation."""
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    
    # Define color scheme (cool colors - blues and teals)
    colors = {
        'primary': RGBColor(0, 102, 204),      # Deep Blue
        'secondary': RGBColor(0, 153, 153),    # Teal
        'accent': RGBColor(102, 178, 255),     # Light Blue
        'dark': RGBColor(25, 25, 112),         # Midnight Blue
        'light': RGBColor(230, 242, 255),      # Very Light Blue
        'white': RGBColor(255, 255, 255),
        'text': RGBColor(51, 51, 51)           # Dark Gray
    }
    
    def add_title_slide(title, subtitle=""):
        """Add a title slide with professional styling."""
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
        
        # Background
        background = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
        )
        background.fill.solid()
        background.fill.fore_color.rgb = colors['primary']
        background.line.fill.background()
        
        # Title
        title_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(2.5), Inches(9), Inches(1.5)
        )
        title_frame = title_box.text_frame
        title_frame.text = title
        title_frame.paragraphs[0].font.size = Pt(54)
        title_frame.paragraphs[0].font.bold = True
        title_frame.paragraphs[0].font.color.rgb = colors['white']
        title_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        
        # Subtitle
        if subtitle:
            subtitle_box = slide.shapes.add_textbox(
                Inches(0.5), Inches(4.5), Inches(9), Inches(1)
            )
            subtitle_frame = subtitle_box.text_frame
            subtitle_frame.text = subtitle
            subtitle_frame.paragraphs[0].font.size = Pt(24)
            subtitle_frame.paragraphs[0].font.color.rgb = colors['accent']
            subtitle_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        
        return slide
    
    def add_section_header(title):
        """Add a section header slide."""
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        
        # Background gradient effect with shapes
        bg1 = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
        )
        bg1.fill.solid()
        bg1.fill.fore_color.rgb = colors['dark']
        bg1.line.fill.background()
        
        # Accent bar
        accent_bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, Inches(3), prs.slide_width, Inches(1.5)
        )
        accent_bar.fill.solid()
        accent_bar.fill.fore_color.rgb = colors['secondary']
        accent_bar.line.fill.background()
        
        # Title
        title_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(3), Inches(9), Inches(1.5)
        )
        title_frame = title_box.text_frame
        title_frame.text = title
        title_frame.paragraphs[0].font.size = Pt(44)
        title_frame.paragraphs[0].font.bold = True
        title_frame.paragraphs[0].font.color.rgb = colors['white']
        title_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        title_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        
        return slide
    
    def add_content_slide(slide_number, title, content_points):
        """Add a content slide with bullet points."""
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        
        # Background
        background = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
        )
        background.fill.solid()
        background.fill.fore_color.rgb = colors['white']
        background.line.fill.background()
        
        # Header bar
        header = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1)
        )
        header.fill.solid()
        header.fill.fore_color.rgb = colors['primary']
        header.line.fill.background()
        
        # Slide number
        slide_num_box = slide.shapes.add_textbox(
            Inches(0.3), Inches(0.25), Inches(0.8), Inches(0.5)
        )
        slide_num_frame = slide_num_box.text_frame
        slide_num_frame.text = f"{slide_number}"
        slide_num_frame.paragraphs[0].font.size = Pt(18)
        slide_num_frame.paragraphs[0].font.bold = True
        slide_num_frame.paragraphs[0].font.color.rgb = colors['white']
        
        # Title
        title_box = slide.shapes.add_textbox(
            Inches(1.2), Inches(0.15), Inches(8.5), Inches(0.7)
        )
        title_frame = title_box.text_frame
        title_frame.text = title
        title_frame.paragraphs[0].font.size = Pt(32)
        title_frame.paragraphs[0].font.bold = True
        title_frame.paragraphs[0].font.color.rgb = colors['white']
        title_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        
        # Content
        content_box = slide.shapes.add_textbox(
            Inches(0.7), Inches(1.5), Inches(8.6), Inches(5.5)
        )
        text_frame = content_box.text_frame
        text_frame.word_wrap = True
        
        for i, point in enumerate(content_points):
            if i > 0:
                p = text_frame.add_paragraph()
            else:
                p = text_frame.paragraphs[0]
            
            p.text = point
            p.level = 0
            p.font.size = Pt(18)
            p.font.color.rgb = colors['text']
            p.space_before = Pt(12)
            p.space_after = Pt(6)
        
        return slide
    
    def add_two_column_slide(slide_number, title, left_content, right_content):
        """Add a two-column content slide."""
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        
        # Background
        background = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
        )
        background.fill.solid()
        background.fill.fore_color.rgb = colors['white']
        background.line.fill.background()
        
        # Header bar
        header = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1)
        )
        header.fill.solid()
        header.fill.fore_color.rgb = colors['primary']
        header.line.fill.background()
        
        # Slide number
        slide_num_box = slide.shapes.add_textbox(
            Inches(0.3), Inches(0.25), Inches(0.8), Inches(0.5)
        )
        slide_num_frame = slide_num_box.text_frame
        slide_num_frame.text = f"{slide_number}"
        slide_num_frame.paragraphs[0].font.size = Pt(18)
        slide_num_frame.paragraphs[0].font.bold = True
        slide_num_frame.paragraphs[0].font.color.rgb = colors['white']
        
        # Title
        title_box = slide.shapes.add_textbox(
            Inches(1.2), Inches(0.15), Inches(8.5), Inches(0.7)
        )
        title_frame = title_box.text_frame
        title_frame.text = title
        title_frame.paragraphs[0].font.size = Pt(32)
        title_frame.paragraphs[0].font.bold = True
        title_frame.paragraphs[0].font.color.rgb = colors['white']
        title_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        
        # Left column
        left_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(1.5), Inches(4.3), Inches(5.5)
        )
        left_frame = left_box.text_frame
        left_frame.word_wrap = True
        for i, point in enumerate(left_content):
            if i > 0:
                p = left_frame.add_paragraph()
            else:
                p = left_frame.paragraphs[0]
            p.text = point
            p.font.size = Pt(16)
            p.font.color.rgb = colors['text']
            p.space_before = Pt(8)
        
        # Right column
        right_box = slide.shapes.add_textbox(
            Inches(5.2), Inches(1.5), Inches(4.3), Inches(5.5)
        )
        right_frame = right_box.text_frame
        right_frame.word_wrap = True
        for i, point in enumerate(right_content):
            if i > 0:
                p = right_frame.add_paragraph()
            else:
                p = right_frame.paragraphs[0]
            p.text = point
            p.font.size = Pt(16)
            p.font.color.rgb = colors['text']
            p.space_before = Pt(8)
        
        return slide
    
    # ========== CREATING 35 SLIDES ==========
    
    # Slide 1: Title Slide
    add_title_slide(
        "Audit Completion & Reporting",
        "A Comprehensive Guide to Professional Audit Practices"
    )
    
    # Slide 2: Overview
    add_content_slide(2, "Overview of Audit Completion & Reporting", [
        "• Understanding the audit lifecycle and completion phase",
        "• Key objectives of audit reporting",
        "• Regulatory and compliance requirements",
        "• Stakeholder communication strategies",
        "• Quality assurance in audit completion"
    ])
    
    # Slide 3: Section Header - Audit Process
    add_section_header("The Audit Process")
    
    # Slide 4: Audit Planning Phase
    add_content_slide(4, "Audit Planning Phase", [
        "• Risk assessment and materiality determination",
        "• Developing audit strategy and plan",
        "• Understanding the entity and its environment",
        "• Identifying areas requiring special audit consideration",
        "• Resource allocation and team assignment"
    ])
    
    # Slide 5: Fieldwork and Evidence Collection
    add_content_slide(5, "Fieldwork and Evidence Collection", [
        "• Substantive testing procedures",
        "• Internal control evaluation",
        "• Sampling methodologies",
        "• Documentation requirements",
        "• Working paper organization and review"
    ])
    
    # Slide 6: Section Header - Audit Completion
    add_section_header("Audit Completion Phase")
    
    # Slide 7: Key Activities in Audit Completion
    add_content_slide(7, "Key Activities in Audit Completion", [
        "• Final analytical procedures",
        "• Evaluation of audit evidence",
        "• Review of subsequent events",
        "• Going concern assessment",
        "• Summary of uncorrected misstatements",
        "• Final materiality considerations"
    ])
    
    # Slide 8: Subsequent Events Review
    add_content_slide(8, "Subsequent Events Review", [
        "• Types of subsequent events (Type I and Type II)",
        "• Review period and procedures",
        "• Management representation requirements",
        "• Disclosure considerations",
        "• Impact on audit opinion"
    ])
    
    # Slide 9: Going Concern Evaluation
    add_content_slide(9, "Going Concern Evaluation", [
        "• Assessment of management's evaluation",
        "• Analysis of financial and operational indicators",
        "• Review of mitigating factors and management plans",
        "• Consideration of post-balance sheet events",
        "• Documentation requirements for going concern"
    ])
    
    # Slide 10: Uncorrected Misstatements
    add_content_slide(10, "Evaluation of Uncorrected Misstatements", [
        "• Accumulation of identified misstatements",
        "• Quantitative and qualitative considerations",
        "• Communication with management and governance",
        "• Impact on audit opinion",
        "• Documentation in summary of uncorrected differences"
    ])
    
    # Slide 11: Section Header - Quality Review
    add_section_header("Quality Review Process")
    
    # Slide 12: Engagement Quality Control Review
    add_content_slide(12, "Engagement Quality Control Review (EQCR)", [
        "• Purpose and scope of EQCR",
        "• Reviewer independence and qualifications",
        "• Key areas of focus in the review",
        "• Documentation of EQCR procedures",
        "• Resolution of matters raised by reviewer"
    ])
    
    # Slide 13: File Review and Archiving
    add_content_slide(13, "File Review and Archiving", [
        "• Internal technical review procedures",
        "• Completion of review points and notes",
        "• File assembly and archiving requirements",
        "• Retention policies and compliance",
        "• Securing audit documentation"
    ])
    
    # Slide 14: Section Header - Management Representations
    add_section_header("Management Representations")
    
    # Slide 15: Written Representations from Management
    add_content_slide(15, "Written Representations from Management", [
        "• Purpose of management representation letter",
        "• Required representations per auditing standards",
        "• Specific representations related to key audit areas",
        "• Timing and dating of representation letter",
        "• Handling refusal to provide representations"
    ])
    
    # Slide 16: Key Representation Areas
    add_two_column_slide(16, "Key Areas Covered in Representations",
        [
            "Financial Statements:",
            "• Completeness and accuracy",
            "• Fair presentation",
            "• Accounting policies",
            "",
            "Internal Controls:",
            "• Design and implementation",
            "• Known deficiencies",
            "• Fraud considerations"
        ],
        [
            "Transactions & Events:",
            "• Completeness of records",
            "• Related party disclosures",
            "• Subsequent events",
            "",
            "Estimates & Judgments:",
            "• Reasonableness of estimates",
            "• Assumptions used",
            "• Disclosure adequacy"
        ]
    )
    
    # Slide 17: Section Header - Audit Reporting
    add_section_header("Audit Reporting Standards")
    
    # Slide 18: Types of Audit Opinions
    add_content_slide(18, "Types of Audit Opinions", [
        "• Unmodified (Unqualified) Opinion",
        "• Modified Opinions:",
        "  - Qualified Opinion",
        "  - Adverse Opinion",
        "  - Disclaimer of Opinion",
        "• Emphasis of Matter and Other Matter paragraphs"
    ])
    
    # Slide 19: Unmodified Audit Opinion
    add_content_slide(19, "Unmodified (Clean) Audit Opinion", [
        "• Financial statements present fairly in all material respects",
        "• Conformity with applicable financial reporting framework",
        "• Sufficient appropriate audit evidence obtained",
        "• Standard report structure and wording",
        "• Appropriate for most audit engagements"
    ])
    
    # Slide 20: Modified Audit Opinions
    add_content_slide(20, "Modified Audit Opinions - Overview", [
        "• Qualified Opinion: Material but not pervasive misstatements",
        "• Adverse Opinion: Material and pervasive misstatements",
        "• Disclaimer of Opinion: Unable to obtain sufficient evidence",
        "• Basis for modification paragraph required",
        "• Impact on opinion paragraph structure"
    ])
    
    # Slide 21: Qualified Opinion
    add_content_slide(21, "Qualified Opinion Details", [
        "• When to issue: Material but not pervasive issues",
        "• 'Except for' language in opinion paragraph",
        "• Common scenarios requiring qualification",
        "• Documentation requirements",
        "• Communication with those charged with governance"
    ])
    
    # Slide 22: Adverse and Disclaimer Opinions
    add_two_column_slide(22, "Adverse vs. Disclaimer Opinions",
        [
            "Adverse Opinion:",
            "• Material and pervasive misstatements",
            "• Financial statements are not fairly presented",
            "• Auditor has obtained sufficient evidence",
            "• Rare in practice",
            "• Significant reputational impact"
        ],
        [
            "Disclaimer of Opinion:",
            "• Unable to obtain sufficient evidence",
            "• Scope limitation is material and pervasive",
            "• Auditor cannot express an opinion",
            "• May result from client-imposed restrictions",
            "• Consider withdrawal from engagement"
        ]
    )
    
    # Slide 23: Emphasis of Matter and Other Matter
    add_content_slide(23, "Emphasis of Matter & Other Matter Paragraphs", [
        "• Emphasis of Matter: Draws attention to FS disclosures",
        "• Other Matter: Relevant to users' understanding of audit",
        "• Does not affect the audit opinion",
        "• Appropriate placement in audit report",
        "• Examples: Going concern, significant uncertainties, etc."
    ])
    
    # Slide 24: Section Header - Audit Report Structure
    add_section_header("Audit Report Structure")
    
    # Slide 25: Components of Audit Report
    add_content_slide(25, "Standard Audit Report Components", [
        "• Report Title (indicating independent audit)",
        "• Addressee (shareholders, board, etc.)",
        "• Opinion section",
        "• Basis for Opinion",
        "• Key Audit Matters (for listed entities)",
        "• Responsibilities sections (management and auditor)",
        "• Auditor's signature, date, and address"
    ])
    
    # Slide 26: Key Audit Matters (KAM)
    add_content_slide(26, "Key Audit Matters (KAM)", [
        "• Required for audits of listed entities",
        "• Matters of most significance in the audit",
        "• Selected from matters communicated to governance",
        "• Description of each KAM and audit response",
        "• Enhances communicative value of audit report",
        "• Does not substitute for required disclosures"
    ])
    
    # Slide 27: Auditor Responsibilities
    add_content_slide(27, "Auditor's Responsibilities", [
        "• Obtain reasonable assurance about FS being free from material misstatement",
        "• Exercise professional judgment and skepticism",
        "• Identify and assess risks of material misstatement",
        "• Obtain sufficient appropriate audit evidence",
        "• Evaluate appropriateness of accounting policies",
        "• Conclude on going concern basis of accounting"
    ])
    
    # Slide 28: Management's Responsibilities
    add_content_slide(28, "Management's Responsibilities", [
        "• Preparation and fair presentation of financial statements",
        "• Design and implementation of internal controls",
        "• Assessment of entity's ability to continue as going concern",
        "• Disclosure of going concern uncertainties if applicable",
        "• Providing auditor with access to information"
    ])
    
    # Slide 29: Section Header - Communications
    add_section_header("Communication with Stakeholders")
    
    # Slide 30: Communication with Those Charged with Governance
    add_content_slide(30, "Communication with Governance", [
        "• Auditor's responsibilities and scope of audit",
        "• Significant findings from the audit",
        "• Internal control deficiencies",
        "• Qualitative aspects of accounting practices",
        "• Auditor independence matters",
        "• Difficulties encountered during audit"
    ])
    
    # Slide 31: Management Letter
    add_content_slide(31, "Management Letter (Letter of Recommendations)", [
        "• Purpose: Communicate control deficiencies and improvement opportunities",
        "• Distinction between significant deficiencies and other matters",
        "• Constructive tone with practical recommendations",
        "• Management response and action plan",
        "• Follow-up on prior year recommendations"
    ])
    
    # Slide 32: Section Header - Best Practices
    add_section_header("Best Practices & Quality")
    
    # Slide 33: Best Practices in Audit Completion
    add_content_slide(33, "Best Practices for Audit Completion", [
        "• Maintain professional skepticism throughout",
        "• Comprehensive documentation and review",
        "• Effective communication with client",
        "• Timely resolution of outstanding matters",
        "• Thorough quality control reviews",
        "• Continuous learning and improvement"
    ])
    
    # Slide 34: Common Challenges and Solutions
    add_two_column_slide(34, "Common Challenges & Solutions",
        [
            "Challenges:",
            "• Time pressures and deadlines",
            "• Complex accounting estimates",
            "• Management bias in judgments",
            "• Incomplete information",
            "• Scope limitations",
            "• Changes in circumstances"
        ],
        [
            "Solutions:",
            "• Effective planning and resource management",
            "• Engage specialists when needed",
            "• Apply professional skepticism",
            "• Clear communication protocols",
            "• Document thoroughly",
            "• Regular progress monitoring"
        ]
    )
    
    # Slide 35: Conclusion
    add_content_slide(35, "Conclusion", [
        "• Audit completion and reporting require careful attention to detail",
        "• Compliance with professional standards is essential",
        "• Quality and independence must be maintained",
        "• Clear communication enhances audit value",
        "• Continuous improvement drives professional excellence",
        "",
        "Thank you for your attention!"
    ])
    
    return prs


def main():
    """Main function to create and save the presentation."""
    print("Creating Audit Completion & Reporting presentation...")
    prs = create_presentation()
    
    output_file = "/home/runner/work/issamfaume/issamfaume/Audit_Completion_and_Reporting.pptx"
    prs.save(output_file)
    print(f"Presentation created successfully: {output_file}")
    print(f"Total slides: {len(prs.slides)}")


if __name__ == "__main__":
    main()
